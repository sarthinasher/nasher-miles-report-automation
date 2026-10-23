"""
2026-10-23: Sync Sales to SharePoint - catch-up of missed days, content guard, Graph
retries / token handling, safe folder creation.
Record: "Errors and their fixes/2026-10-23-sharepoint-sales-sync-no-catch-up.md".
"""
import importlib.util
import os
import sys
from datetime import date, datetime, timedelta

import pytest
import requests

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from common import sharepoint_utils as sp  # noqa: E402


def _load_sync():
    path = os.path.join(os.path.dirname(__file__), "..", "scripts", "sync_sales_to_sharepoint.py")
    spec = importlib.util.spec_from_file_location("sync_sales_to_sharepoint", path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


sync = _load_sync()
D = lambda d, m=10: date(2026, m, d)  # noqa: E731
HEADER = ["Channel", "Order No", "ExtOrderNo", "Order Date", "Order Type", "status",
          "Customer Name", "SKU Code", "Quantity", "Invoice No", "Invoice Date", "Total"]


def row(day, n=1):
    return ["online", f"P{n}", f"X{n}", f"{day:%Y-%m-%d} 01:00:00", "Prepaid", "Shipped",
            "C", "SKU", 1, f"INV{n}", f"{day:%Y-%m-%d} 14:30:28", 2199]


# ── which dates ──────────────────────────────────────────────
def test_blank_report_date_is_a_7_day_catch_up_not_just_yesterday():
    days, explicit, mode = sync.sync_dates(D(23), "")
    assert days == [D(16), D(17), D(18), D(19), D(20), D(21), D(22)]
    assert explicit is False and "catch-up" in mode


def test_requested_dates_are_explicit():
    days, explicit, _ = sync.sync_dates(D(23), "11/10/2026-13/10/2026")
    assert days == [D(11), D(12), D(13)] and explicit is True


def test_requested_future_date_rejected():
    from common.report_dates import ReportDateError
    with pytest.raises(ReportDateError):
        sync.sync_dates(D(23), "23/10/2026")


DRIVE = {"id": "1", "name": "21.10.2026", "modifiedTime": "2026-10-23T06:24:22.535Z"}


def test_missing_on_sharepoint_is_synced():
    assert sync.sync_reason(DRIVE, None, False)[0] is True


def test_reexported_in_drive_after_sharepoint_copy_is_synced_again():
    old = {"lastModifiedDateTime": "2026-10-12T11:48:40Z"}
    up, reason = sync.sync_reason(DRIVE, old, False)
    assert up is True and "re-exported" in reason


def test_up_to_date_copy_is_skipped_unless_requested():
    newer = {"lastModifiedDateTime": "2026-10-23T09:14:38Z"}
    assert sync.sync_reason(DRIVE, newer, False)[0] is False
    assert sync.sync_reason(DRIVE, newer, True) == (True, "requested")


def test_missing_in_drive_is_not_synced():
    assert sync.sync_reason(None, None, False)[0] is False
    assert sync.sync_reason(None, None, True)[0] is False
    assert "already has it" in sync.sync_reason(None, {"lastModifiedDateTime": "x"}, False)[1]


# ── content guard ────────────────────────────────────────────
def test_one_column_file_is_refused():
    # The broken Sales files of 20/21-Oct: header "Order No" only, 0 rows.
    with pytest.raises(sync.SalesFileError, match="only 1 column"):
        sync.check_sales_rows([["Order No"]], D(14))


def test_empty_file_is_refused():
    with pytest.raises(sync.SalesFileError, match="empty"):
        sync.check_sales_rows([], D(14))
    with pytest.raises(sync.SalesFileError, match="empty"):
        sync.check_sales_rows([[None, ""], [None]], D(14))


def test_header_only_file_is_refused():
    with pytest.raises(sync.SalesFileError, match="0 data rows"):
        sync.check_sales_rows([HEADER, [None] * len(HEADER)], D(14))


def test_missing_required_column_is_refused():
    header = [h for h in HEADER if h != "Invoice Date"] + ["Extra"]
    with pytest.raises(sync.SalesFileError, match="Invoice Date"):
        sync.check_sales_rows([header, ["x"] * len(header)], D(14))


def test_file_of_another_day_is_refused():
    # 2026-10-12: Drive 10.10.2026 had been uploaded as 11.10.2026.xlsx.
    with pytest.raises(sync.SalesFileError, match="another day"):
        sync.check_sales_rows([HEADER, row(D(10)), row(D(10), 2)], D(11))


def test_good_file_passes_with_string_or_datetime_invoice_dates():
    assert sync.check_sales_rows([HEADER, row(D(22)), row(D(22), 2)], D(22)) == (2, len(HEADER))
    r = row(D(22))
    r[HEADER.index("Invoice Date")] = datetime(2026, 10, 22, 14, 30)
    assert sync.check_sales_rows([HEADER, r, row(D(21), 2)], D(22)) == (2, len(HEADER))
    r2 = row(D(22))
    r2[HEADER.index("Invoice Date")] = "unparseable"
    assert sync.check_sales_rows([HEADER, r2], D(22)) == (1, len(HEADER))


def _xlsx(tmp_path, rows, title="Sales Report", first=None):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    if first:
        ws.title = first
        ws = wb.create_sheet(title)
    else:
        ws.title = title
    for r in rows:
        ws.append(r)
    p = tmp_path / "22.10.2026.xlsx"
    wb.save(p)
    return str(p)


def test_validate_sales_file_reads_the_sales_report_sheet(tmp_path):
    p = _xlsx(tmp_path, [HEADER, row(D(22)), row(D(22), 2), row(D(22), 3)], first="Other")
    assert sync.validate_sales_file(p, D(22)) == (3, len(HEADER))


def test_validate_sales_file_refuses_junk(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_bytes(b"")
    with pytest.raises(sync.SalesFileError, match="0 bytes"):
        sync.validate_sales_file(str(p), D(22))
    p.write_bytes(b"not a zip")
    with pytest.raises(sync.SalesFileError, match="readable"):
        sync.validate_sales_file(str(p), D(22))


# ── main(): the 12.10-22.10 situation end to end (all I/O faked) ─────
def _run_main(monkeypatch, tmp_path, today, drive, sp_files, report_date="", broken=()):
    uploads = []

    class FakeDT(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(today.year, today.month, today.day, 12, 0, tzinfo=tz)

    for k in ("GDRIVE_SERVICE_ACCOUNT_JSON", "SP_CLIENT_ID", "SP_TENANT_ID", "SP_CLIENT_SECRET"):
        monkeypatch.setenv(k, "x")
    monkeypatch.setenv("REPORT_DATE", report_date)
    monkeypatch.setattr(sync, "datetime", FakeDT)
    monkeypatch.setattr(sync, "LOCAL_TMP", str(tmp_path))
    monkeypatch.setattr(sync, "get_drive_service", lambda: object())
    monkeypatch.setattr(sync, "find_drive_file_by_name", lambda svc, folder, name: drive.get(name))
    monkeypatch.setattr(sync, "GraphAuth", lambda: "auth")
    monkeypatch.setattr(sync, "resolve_site_id", lambda t, u: "site")
    monkeypatch.setattr(sync, "resolve_drive_id", lambda t, s: "drive")
    monkeypatch.setattr(sync, "ensure_folder_path", lambda t, d, p: p)
    monkeypatch.setattr(sync, "list_folder_files", lambda t, d, p: sp_files)

    def download(svc, meta, dest):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales Report"
        day = datetime.strptime(meta["name"], "%d.%m.%Y").date()
        for r in ([["Order No"]] if meta["name"] in broken else [HEADER, row(day)]):
            ws.append(r)
        wb.save(dest)
        return dest

    monkeypatch.setattr(sync, "download_drive_file_as_xlsx", download)
    monkeypatch.setattr(sync, "upload_file_to_sharepoint",
                        lambda t, d, local, folder, filename: uploads.append(filename) or "url")
    code = 0
    try:
        sync.main()
    except SystemExit as e:
        code = e.code
    return code, uploads


def _meta(name, modified="2026-10-23T06:20:00Z"):
    return {"id": name, "name": name, "mimeType": "application/vnd.google-apps.spreadsheet",
            "modifiedTime": modified}


def test_22_oct_run_old_behaviour_now_fails_only_for_the_missing_day(monkeypatch, tmp_path):
    # 22-Oct: 21.10 not in Drive yet; 19.10 was (created 21-Oct 15:50) but never mirrored.
    drive = {"19.10.2026": _meta("19.10.2026", "2026-10-21T15:50:50Z")}
    sp_files = {f"{d:%d.%m.%Y}.xlsx": {"lastModifiedDateTime": "2026-10-12T11:48:40Z"}
                for d in [D(11)]}
    code, uploads = _run_main(monkeypatch, tmp_path, D(22), drive, sp_files)
    assert uploads == ["19.10.2026.xlsx"]          # caught up instead of ignored
    assert code == 1                                # 15-18, 20, 21 still missing


def test_23_oct_catch_up_mirrors_every_missing_day(monkeypatch, tmp_path):
    drive = {f"{d:%d.%m.%Y}": _meta(f"{d:%d.%m.%Y}") for d in [D(i) for i in range(11, 23)]}
    sp_files = {"22.10.2026.xlsx": {"lastModifiedDateTime": "2026-10-23T09:14:38Z"}}
    code, uploads = _run_main(monkeypatch, tmp_path, D(23), drive, sp_files)
    assert uploads == [f"{i}.10.2026.xlsx" for i in range(16, 22)]   # 22.10 up to date
    assert code == 0


def test_broken_file_is_not_mirrored_and_fails_the_run(monkeypatch, tmp_path):
    drive = {f"{d:%d.%m.%Y}": _meta(f"{d:%d.%m.%Y}") for d in [D(i) for i in range(14, 17)]}
    code, uploads = _run_main(monkeypatch, tmp_path, D(23), drive, {},
                              report_date="14/10/2026-16/10/2026", broken={"15.10.2026"})
    assert uploads == ["14.10.2026.xlsx", "16.10.2026.xlsx"]
    assert code == 1


def test_requested_dates_reupload_even_if_up_to_date(monkeypatch, tmp_path):
    drive = {"11.10.2026": _meta("11.10.2026")}
    sp_files = {"11.10.2026.xlsx": {"lastModifiedDateTime": "2026-10-23T10:00:00Z"}}
    code, uploads = _run_main(monkeypatch, tmp_path, D(23), drive, sp_files,
                              report_date="11/10/2026")
    assert (code, uploads) == (0, ["11.10.2026.xlsx"])


# ── Graph: retries, token, folders, upload ───────────────────
class Resp:
    def __init__(self, status, body=None, headers=None, text=""):
        self.status_code, self._body, self.headers, self.text = status, body, headers or {}, text

    def json(self):
        if self._body is None:
            raise ValueError("no json")
        return self._body


@pytest.fixture
def graph(monkeypatch):
    """Scripted Graph: graph.queue = [Resp | Exception, ...]; graph.calls = [(method, url, kw)]."""
    class G:
        queue, calls, sleeps = [], [], []

    g = G()

    def send(method, url, **kw):
        g.calls.append((method, url, kw))
        item = g.queue.pop(0)
        if isinstance(item, Exception):
            raise item
        return item

    monkeypatch.setattr(sp, "_send", send)
    monkeypatch.setattr(sp, "_sleep", g.sleeps.append)
    monkeypatch.setattr(sp, "_rand", lambda: 0.0)
    return g


def test_graph_retries_503_and_dropped_connection(graph):
    graph.queue = [Resp(503), requests.exceptions.ConnectionError("reset"), Resp(200, {"id": "S"})]
    assert sp.resolve_site_id("tok", "https://x.sharepoint.com/sites/NM") == "S"
    assert len(graph.calls) == 3 and graph.sleeps == [2.0, 4.0]
    assert graph.calls[0][2]["headers"]["Authorization"] == "Bearer tok"


def test_graph_honours_retry_after_on_429(graph):
    graph.queue = [Resp(429, headers={"Retry-After": "7"}), Resp(200, {"id": "D"})]
    assert sp.resolve_drive_id("tok", "S") == "D"
    assert graph.sleeps == [7.0]


def test_graph_gives_up_after_max_attempts(graph):
    graph.queue = [Resp(503)] * sp.MAX_ATTEMPTS
    with pytest.raises(sp.SharePointError, match="503"):
        sp.resolve_drive_id("tok", "S")
    assert len(graph.calls) == sp.MAX_ATTEMPTS


def test_graph_does_not_retry_403(graph):
    graph.queue = [Resp(403, {"error": {"code": "accessDenied", "message": "no"}})]
    with pytest.raises(sp.SharePointError, match="accessDenied"):
        sp.resolve_drive_id("tok", "S")
    assert len(graph.calls) == 1


def test_connection_errors_after_retries_hide_upload_url_query(graph):
    err = requests.exceptions.ConnectionError("url: /up?guid=1&tempauth=SECRET")
    graph.queue = [err] * sp.MAX_ATTEMPTS
    with pytest.raises(sp.SharePointError) as e:
        sp.graph_request("PUT", "https://t.sharepoint.com/up?tempauth=SECRET", None)
    assert "SECRET" not in str(e.value)


def _token_ok(tok="T", expires_in=3600):
    return Resp(200, {"access_token": tok, "expires_in": expires_in})


@pytest.fixture
def creds(monkeypatch):
    for k in ("SP_CLIENT_ID", "SP_TENANT_ID", "SP_CLIENT_SECRET"):
        monkeypatch.setenv(k, k.lower())


def test_401_renews_the_token_once(graph, creds):
    auth = sp.GraphAuth()
    graph.queue = [_token_ok("T1"), Resp(401), _token_ok("T2"), Resp(200, {"id": "D"})]
    assert sp.resolve_drive_id(auth, "S") == "D"
    assert graph.calls[1][2]["headers"]["Authorization"] == "Bearer T1"
    assert graph.calls[3][2]["headers"]["Authorization"] == "Bearer T2"
    assert graph.sleeps == []


def test_token_renewed_before_it_expires(graph, creds, monkeypatch):
    clock = [1000.0]
    monkeypatch.setattr(sp, "_now", lambda: clock[0])
    auth = sp.GraphAuth()
    graph.queue = [_token_ok("T1", 3600), _token_ok("T2", 3600)]
    assert auth.token() == "T1"
    clock[0] += 3000                                 # 10 min left: still fine
    assert auth.token() == "T1"
    clock[0] += 400                                  # < 5 min left: renew
    assert auth.token() == "T2"


def test_expired_secret_error_is_explained(graph, creds):
    graph.queue = [Resp(401, {"error": "invalid_client", "error_codes": [7000222],
                              "error_description": "AADSTS7000222: The provided client secret "
                                                   "keys are expired.\r\nTrace ID: x"})]
    with pytest.raises(sp.SharePointAuthError) as e:
        sp.get_graph_token()
    assert "AADSTS7000222" in str(e.value) and "EXPIRED" in str(e.value)
    assert "sp_client_secret" not in str(e.value)    # the secret value is never shown
    assert len(graph.calls) == 1                     # not retried


def test_token_request_retries_5xx(graph, creds):
    graph.queue = [Resp(503), _token_ok("T")]
    assert sp.get_graph_token() == "T"


def test_missing_credentials(monkeypatch):
    monkeypatch.delenv("SP_CLIENT_SECRET", raising=False)
    with pytest.raises(sp.SharePointAuthError, match="SP_CLIENT_SECRET"):
        sp.fetch_token()


def test_ensure_folder_never_replaces_an_existing_folder(graph):
    # Before: any non-200 lookup (e.g. a 503) led to a POST with conflictBehavior=replace.
    graph.queue = [Resp(503)] * sp.MAX_ATTEMPTS
    with pytest.raises(sp.SharePointError):
        sp.ensure_folder_path("tok", "D", "Sales/Vinculum")
    assert all(m == "GET" for m, _, _ in graph.calls)


def test_ensure_folder_creates_only_after_404_with_fail(graph):
    graph.queue = [Resp(200, {}), Resp(404), Resp(409)]
    assert sp.ensure_folder_path("tok", "D", "Sales/Vinculum") == "Sales/Vinculum"
    method, url, kw = graph.calls[2]
    assert method == "POST" and url.endswith("root:/Sales:/children")
    assert kw["json"]["@microsoft.graph.conflictBehavior"] == "fail"


def test_list_folder_files_pages_and_404(graph):
    graph.queue = [Resp(200, {"value": [{"name": "a.xlsx", "size": 1, "file": {},
                                         "lastModifiedDateTime": "t1"},
                                        {"name": "sub", "folder": {}}],
                              "@odata.nextLink": "https://graph.microsoft.com/v1.0/next"}),
                   Resp(200, {"value": [{"name": "b.xlsx", "size": 2, "file": {},
                                         "lastModifiedDateTime": "t2"}]})]
    files = sp.list_folder_files("tok", "D", "Sales/Vinculum")
    assert sorted(files) == ["a.xlsx", "b.xlsx"] and files["b.xlsx"]["size"] == 2
    graph.queue = [Resp(404)]
    assert sp.list_folder_files("tok", "D", "Nope") == {}


def test_simple_upload_retried_and_checked(graph, tmp_path):
    p = tmp_path / "22.10.2026.xlsx"
    p.write_bytes(b"x" * 100)
    graph.queue = [Resp(502), Resp(200, {"webUrl": "W", "size": 100})]
    assert sp.upload_file_to_sharepoint("tok", "D", str(p), "Sales/Vinculum",
                                        filename="22.10.2026.xlsx") == "W"
    assert graph.calls[1][1].endswith("root:/Sales/Vinculum/22.10.2026.xlsx:/content")
    graph.queue = [Resp(200, {"webUrl": "W", "size": 0})]
    with pytest.raises(sp.SharePointError, match="0 bytes"):
        sp.upload_file_to_sharepoint("tok", "D", str(p), "Sales/Vinculum")


def test_large_upload_chunks_without_auth_header(graph, tmp_path, monkeypatch):
    monkeypatch.setattr(sp, "SIMPLE_UPLOAD_LIMIT", 10)
    monkeypatch.setattr(sp, "UPLOAD_CHUNK_SIZE", 8)
    p = tmp_path / "big.xlsx"
    p.write_bytes(b"y" * 20)
    graph.queue = [Resp(200, {"uploadUrl": "https://t/up?tempauth=z"}),
                   Resp(202, {}), Resp(503), Resp(202, {}), Resp(201, {"webUrl": "W", "size": 20})]
    assert sp.upload_file_to_sharepoint("tok", "D", str(p), "F") == "W"
    chunk_calls = graph.calls[1:]
    assert all("Authorization" not in kw["headers"] for _, _, kw in chunk_calls)
    assert [kw["headers"]["Content-Range"] for _, _, kw in chunk_calls] == [
        "bytes 0-7/20", "bytes 8-15/20", "bytes 8-15/20", "bytes 16-19/20"]
