"""
Nasher Miles - Drive -> SharePoint sync (Vinculum Sales)
==========================================================
Mirrors the Vinculum Sales daily files (Google Drive folder, named DD.MM.YYYY = invoice
date) into the SharePoint 'Sales/Vinculum' folder as DD.MM.YYYY.xlsx (replacing a
same-named file). The file is always picked BY NAME, never "the newest file"
(2026-10-13: 10-Oct invoices had been uploaded as 11.10.2026.xlsx).

Which dates (2026-10-23):
  - REPORT_DATE blank (scheduled runs) = catch-up: the 7 days ending yesterday (IST).
    A day is synced when it is not on SharePoint yet, or when its Drive file was
    re-exported after the SharePoint copy was written. Before, only yesterday was
    looked at, so 12.10-21.10 were never mirrored: the Sales export for each day
    landed after that day's sync runs had given up.
  - REPORT_DATE=DD/MM/YYYY[,..|-..] = exactly those days, always re-uploaded.

Every file is checked before upload (the Sales sheet must have a header of at least
MIN_COLUMNS named columns incl. REQUIRED_COLUMNS, at least one data row, and invoice
dates of that day): an empty / one-column / wrong-day file is never mirrored over
SharePoint. The run exits 1 if any day could not be synced.

Auth:
  - Google Drive: service account (GDRIVE_SERVICE_ACCOUNT_JSON)  [read side]
  - SharePoint  : Azure AD app-only (SP_CLIENT_ID / SP_TENANT_ID / SP_CLIENT_SECRET)

Run locally for testing:
    GDRIVE_SERVICE_ACCOUNT_JSON='{...}' \
    SP_CLIENT_ID=... SP_TENANT_ID=... SP_CLIENT_SECRET=... \
        python scripts/sync_sales_to_sharepoint.py
"""

import sys
import os
from collections import Counter
from datetime import date, datetime, timedelta, timezone

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from common.drive_utils import (
    get_drive_service,
    find_drive_file_by_name,
    download_drive_file_as_xlsx,
    with_drive_retries,
)
from common.report_dates import (
    DEFAULT_LOOKBACK_DAYS, SALES_REPORT_NAME, ReportDateError, default_window,
    drive_filename, parse_report_dates,
)
from common.sharepoint_utils import (
    GraphAuth,
    resolve_site_id,
    resolve_drive_id,
    ensure_folder_path,
    list_folder_files,
    upload_file_to_sharepoint,
)

# ── SETTINGS ──────────────────────────────────────────────
# Source: Vinculum Sales Drive folder (same one the Sales report writes to)
DRIVE_FOLDER_ID = os.environ.get(
    "VINCULUM_SALES_DRIVE_FOLDER_ID", "1iL94fl0nUOxZ_pEWH_AuxnqBYsYjsjt_"
)

# Destination: SharePoint site + folder path within its document library
SHAREPOINT_SITE_URL = os.environ.get(
    "SP_SITE_URL", "https://nashermilesprivatelimite452.sharepoint.com/sites/NasherMiles"
)
SHAREPOINT_FOLDER_PATH = os.environ.get("SP_SALES_FOLDER", "Sales/Vinculum")

LOCAL_TMP = "/tmp/sp_sync"
IST = timezone(timedelta(hours=5, minutes=30))   # India: no DST

# Content guard (2026-10-23). A Vinculum MIS Sales sheet has ~110 columns; the broken
# files of 20/21-Oct had ONE ("Order No") and no rows.
MIN_COLUMNS = 10
REQUIRED_COLUMNS = ("Order No", "Invoice Date")
INVOICE_DATE_COLUMN = "Invoice Date"


class SalesFileError(ValueError):
    """The Drive file is not a usable Sales report for its date - not mirrored."""


def fail(message, code=1):
    print(f"\n[✗] {message}")
    sys.exit(code)


def validate_env():
    missing = [n for n, v in [
        ("GDRIVE_SERVICE_ACCOUNT_JSON", os.environ.get("GDRIVE_SERVICE_ACCOUNT_JSON")),
        ("SP_CLIENT_ID", os.environ.get("SP_CLIENT_ID")),
        ("SP_TENANT_ID", os.environ.get("SP_TENANT_ID")),
        ("SP_CLIENT_SECRET", os.environ.get("SP_CLIENT_SECRET")),
    ] if not v]
    if missing:
        fail(f"Missing required environment variables/secrets: {', '.join(missing)}")


# ── which dates ──────────────────────────────────────────────
def sync_dates(today, spec, lookback_days=DEFAULT_LOOKBACK_DAYS):
    """(dates, explicit, mode text). REPORT_DATE given -> exactly those dates
    (explicit=True: always re-uploaded); blank -> the lookback_days days ending
    yesterday (IST), oldest first (explicit=False: only missing / stale ones)."""
    requested = parse_report_dates(spec, today)
    if requested:
        return requested, True, f"requested dates ({spec.strip()}) - re-uploaded"
    return (default_window(today, lookback_days), False,
            f"catch-up (last {lookback_days} days ending yesterday, IST: "
            "missing on SharePoint or re-exported in Drive since)")


def parse_ts(text):
    """Drive / Graph RFC 3339 time ('2026-10-23T06:26:19.467Z') -> aware datetime."""
    if not text:
        return None
    try:
        return datetime.fromisoformat(str(text).replace("Z", "+00:00"))
    except ValueError:
        return None


def sync_reason(drive_meta, sp_item, explicit):
    """Why a day is (or is not) uploaded: (upload?, reason)."""
    if drive_meta is None:
        if sp_item is not None:
            return False, "not in Drive; SharePoint already has it"
        return False, "not in Drive (Vinculum Sales export missing)"
    if explicit:
        return True, "requested"
    if sp_item is None:
        return True, "not on SharePoint yet"
    drive_ts = parse_ts(drive_meta.get("modifiedTime"))
    sp_ts = parse_ts(sp_item.get("lastModifiedDateTime"))
    if drive_ts and sp_ts and drive_ts > sp_ts:
        return True, (f"Drive file re-exported ({drive_meta.get('modifiedTime')}) after the "
                      f"SharePoint copy ({sp_item.get('lastModifiedDateTime')})")
    return False, "SharePoint copy is up to date"


# ── content guard ────────────────────────────────────────────
def _blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _cell_date(v):
    """date of an Invoice Date cell (datetime / date / '2026-10-22 14:30:28' /
    '22/10/2026 ...' / '22-10-2026'), or None."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if not isinstance(v, str):
        return None
    text = v.strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def check_sales_rows(rows, day):
    """(data rows, named columns) of a Sales sheet, or SalesFileError when it is not
    a usable Sales report for `day`: no / too few columns, a required column missing,
    no data row, or no invoice dated `day` (a file of another day under this name)."""
    rows = [list(r) for r in rows or []]
    while rows and all(_blank(v) for v in rows[0]):
        rows.pop(0)
    if not rows:
        raise SalesFileError("the file is empty (no header, no rows)")
    header = ["" if h is None else str(h).strip() for h in rows[0]]
    named = [h for h in header if h]
    if len(named) < MIN_COLUMNS:
        raise SalesFileError(f"only {len(named)} column(s) {named[:5]} - a Sales report has "
                             f"~110; the file is broken")
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise SalesFileError(f"column(s) {missing} missing (has {len(named)} columns)")
    data = [r for r in rows[1:] if not all(_blank(v) for v in r)]
    if not data:
        raise SalesFileError(f"header only, 0 data rows ({len(named)} columns)")
    col = header.index(INVOICE_DATE_COLUMN)
    dates = Counter(_cell_date(r[col]) for r in data if col < len(r))
    dates.pop(None, None)
    if dates and dates.get(day, 0) == 0:
        top = ", ".join(f"{d.strftime('%d.%m.%Y')} x{n}" for d, n in dates.most_common(3))
        raise SalesFileError(f"no invoice dated {day.strftime('%d.%m.%Y')} in {len(data)} rows "
                             f"(invoice dates: {top}) - the file is for another day")
    return len(data), len(named)


def read_sales_rows(path):
    """Rows of the Sales sheet ('Sales Report' if present, else the first sheet).
    reset_dimensions(): Vinculum xlsx can declare a wrong size (2026-10-21)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = next((w for w in wb.worksheets if w.title.strip().lower() == "sales report"),
                  wb.worksheets[0])
        ws.reset_dimensions()
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def validate_sales_file(path, day):
    if os.path.getsize(path) == 0:
        raise SalesFileError("the downloaded file is 0 bytes")
    try:
        rows = read_sales_rows(path)
    except SalesFileError:
        raise
    except Exception as e:
        raise SalesFileError(f"not a readable .xlsx ({type(e).__name__}: {e})") from e
    return check_sales_rows(rows, day)


# ── one day ──────────────────────────────────────────────────
def sync_day(day, meta, drive_service, sp_auth, drive_id):
    """Downloads, checks and uploads one day. Returns the webUrl; raises on any problem."""
    xlsx_name = drive_filename(day, *SALES_REPORT_NAME) + ".xlsx"
    local_path = os.path.join(LOCAL_TMP, xlsx_name)
    try:
        with_drive_retries(lambda svc: download_drive_file_as_xlsx(svc, meta, local_path),
                           drive_service, f"download of {meta['name']}")
        n_rows, n_cols = validate_sales_file(local_path, day)
        print(f"    [OK] {xlsx_name}: {n_rows} rows x {n_cols} columns, invoices of "
              f"{day.strftime('%d.%m.%Y')} ({os.path.getsize(local_path)} bytes)")
        return upload_file_to_sharepoint(sp_auth, drive_id, local_path, SHAREPOINT_FOLDER_PATH,
                                         filename=xlsx_name)
    finally:
        try:
            os.remove(local_path)
        except OSError:
            pass


def main():
    validate_env()
    os.makedirs(LOCAL_TMP, exist_ok=True)

    print("=" * 55)
    print("  Nasher Miles - Drive -> SharePoint (Vinculum Sales)")
    print("=" * 55)

    today = datetime.now(IST).date()
    try:
        days, explicit, mode = sync_dates(today, os.environ.get("REPORT_DATE", ""))
    except ReportDateError as e:
        fail(f"REPORT_DATE: {e}", code=2)
    print(f"\n    Mode: {mode}")
    print(f"    Dates: {', '.join(drive_filename(d, *SALES_REPORT_NAME) for d in days)}")

    print("\n[1] Authenticating Google Drive (service account)...")
    drive_service = get_drive_service()
    print("    [OK] Drive ready.")

    print("\n[2] Authenticating SharePoint (Microsoft Graph app-only)...")
    sp_auth = GraphAuth()
    try:
        site_id = resolve_site_id(sp_auth, SHAREPOINT_SITE_URL)
        drive_id = resolve_drive_id(sp_auth, site_id)
        ensure_folder_path(sp_auth, drive_id, SHAREPOINT_FOLDER_PATH)
        sp_files = list_folder_files(sp_auth, drive_id, SHAREPOINT_FOLDER_PATH)
    except Exception as e:
        fail(f"SharePoint: {type(e).__name__}: {e}")
    print(f"    [OK] Site + library resolved; {SHAREPOINT_FOLDER_PATH} has {len(sp_files)} file(s).")

    print(f"\n[3] Checking {len(days)} date(s) in Drive and SharePoint...")
    plan, problems, synced = [], [], []
    for day in days:
        name = drive_filename(day, *SALES_REPORT_NAME)
        try:
            meta = with_drive_retries(
                lambda svc: find_drive_file_by_name(svc, DRIVE_FOLDER_ID, name),
                drive_service, f"lookup of {name}")
        except Exception as e:
            problems.append(f"{name}: Drive lookup failed ({type(e).__name__}: {e})")
            print(f"    [✗] {name}: Drive lookup failed: {e}")
            continue
        upload, reason = sync_reason(meta, sp_files.get(name + ".xlsx"), explicit)
        print(f"    {'[->]' if upload else '[  ]'} {name}: {reason}")
        if upload:
            plan.append((day, meta))
        elif meta is None and name + ".xlsx" not in sp_files:
            problems.append(f"{name}: not in Drive, not on SharePoint")

    for i, (day, meta) in enumerate(plan, 1):
        name = drive_filename(day, *SALES_REPORT_NAME)
        print(f"\n[4.{i}] {meta['name']} (Drive, modified {meta.get('modifiedTime')}) -> "
              f"{SHAREPOINT_FOLDER_PATH}/{name}.xlsx")
        try:
            web_url = sync_day(day, meta, drive_service, sp_auth, drive_id)
        except SalesFileError as e:
            problems.append(f"{name}: NOT mirrored - {e}")
            print(f"    [✗] NOT mirrored: {e}. Re-export it (Vinculum Sales Report to Drive, "
                  f"report_date={day.strftime('%d/%m/%Y')}); SharePoint is left unchanged.")
            continue
        except Exception as e:
            problems.append(f"{name}: upload failed ({type(e).__name__}: {e})")
            print(f"    [✗] Failed: {type(e).__name__}: {e}")
            continue
        synced.append(name)
        print(f"    [OK] Uploaded (replaces a same-named file). webUrl: {web_url}")

    print("\n" + "=" * 55)
    print(f"  Synced: {', '.join(synced) or 'nothing'}")
    print(f"  Site  : {SHAREPOINT_SITE_URL}")
    print(f"  Folder: {SHAREPOINT_FOLDER_PATH}")
    print("=" * 55)
    if problems:
        fail(f"{len(problems)} date(s) not synced:\n    - " + "\n    - ".join(problems)
             + "\n  The next scheduled run retries the last 7 days by itself; older days: "
               "run the workflow with report_date.")


if __name__ == "__main__":
    main()
