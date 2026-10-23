"""
Merge one report's per-location exports into the single file uploaded to Drive
(2026-10-19).

Vinculum exports only the session's location (see common/vinculum_ui.py, "Locations"),
so the report scripts export once per location and merge here. The merged file keeps
the file type (.csv / .xlsx), the first sheet's name, the header (column names and
order) and every row of every location - the Drive file name, tab and columns the
downstream FMS / Apps Script jobs read stay exactly as before. Columns such as Pick
Pack "Source WH" / Sales "Originating warehouse" still say which warehouse a row is
from; nothing is added.

Rules:
  * rows are matched to the first file's header BY COLUMN NAME (a column only some
    files have is appended at the end, never dropped);
  * blank rows are skipped;
  * a row that is identical to a row of an EARLIER location's file is dropped (the
    same order seen from two locations); duplicates inside one file are kept;
  * one file with data -> that file is used unchanged (byte for byte).

Pure functions (merge_tables, count_rows) are unit-tested in tests/test_report_merge.py.
"""

import csv
import io
import os


class MergeError(Exception):
    quiet = True


def _blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _trim(row):
    row = list(row)
    while row and _blank(row[-1]):
        row.pop()
    return row


def _keys(header):
    """(name, occurrence) per column, so repeated / empty names still line up."""
    seen, keys = {}, []
    for h in header:
        name = "" if h is None else str(h).strip()
        n = seen.get(name, 0)
        seen[name] = n + 1
        keys.append((name, n))
    return keys


def merge_tables(tables):
    """tables = [(location, rows)] with rows[0] = header (rows may be [] = no data).
    Returns (header, data_rows, counts {location: data rows kept}, duplicates_dropped)."""
    header = None
    out, counts, dropped = [], {}, 0
    seen_before = set()                  # rows of earlier locations' files
    for loc, rows in tables:
        rows = [_trim(r) for r in (rows or [])]
        while rows and not rows[0]:
            rows.pop(0)
        counts.setdefault(loc, 0)
        if not rows:
            continue
        this_header, data = rows[0], rows[1:]
        if header is None:
            header = list(this_header)
        base_keys = _keys(header)
        index = {k: i for i, k in enumerate(base_keys)}
        mapping = []
        for k, name in zip(_keys(this_header), this_header):
            if k not in index:
                header.append(name)
                index[k] = len(header) - 1
            mapping.append(index[k])
        this_file = set()
        for r in data:
            if all(_blank(v) for v in r):
                continue
            while len(mapping) < len(r):     # cell beyond this file's header: keep it
                header.append("")
                mapping.append(len(header) - 1)
            row = [None] * len(header)
            for j, v in enumerate(r):
                row[mapping[j]] = v
            key = tuple("" if _blank(v) else v for v in _trim(row))
            if key in seen_before:
                dropped += 1
                continue
            this_file.add(key)
            out.append(row)
            counts[loc] += 1
        seen_before |= this_file
    width = len(header or [])
    out = [r + [None] * (width - len(r)) for r in out]
    return header or [], out, counts, dropped


# ── file I/O ─────────────────────────────────────────────────
def _decode(raw):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8"


def read_csv_rows(path):
    with open(path, "rb") as f:
        text, _ = _decode(f.read())
    return [list(r) for r in csv.reader(io.StringIO(text))]


def read_xlsx_rows(path):
    """(sheet title, rows) of the FIRST sheet.
    2026-10-21: Vinculum's Order Enquiry .xlsx declares <dimension ref="A1"/>; read-only
    openpyxl trusts it and returned only cell A1 (merged files had just "Order No", 0 rows).
    reset_dimensions() makes it read every cell that is really there."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        ws.reset_dimensions()
        return ws.title, [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        return None, read_csv_rows(path)
    if ext == ".xlsx":
        return read_xlsx_rows(path)
    raise MergeError(f"Cannot merge {ext} files ({path}).")


def count_rows(rows):
    """Non-blank data rows below the header."""
    rows = [_trim(r) for r in rows or []]
    while rows and not rows[0]:
        rows.pop(0)
    return sum(1 for r in rows[1:] if r and not all(_blank(v) for v in r))


def write_csv(path, header, rows):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(["" if v is None else v for v in header])
        for r in rows:
            w.writerow(["" if v is None else v for v in r])


def write_xlsx(path, sheet_title, header, rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "Sheet1")[:31]
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def merge_location_files(parts, out_path):
    """parts = [(location, local file or None when the location had no data)].
    Writes the merged file to out_path (same extension as the parts) and returns
    (path_to_upload, counts {location: rows}, total_rows, duplicates_dropped).
    path_to_upload is None when no location returned a file."""
    files = [(loc, p) for loc, p in parts if p]
    if not files:
        return None, {loc: 0 for loc, _ in parts}, 0, 0
    exts = {os.path.splitext(p)[1].lower() for _, p in files}
    if len(exts) != 1:
        raise MergeError(f"Per-location files have different types: {sorted(exts)}")
    ext = exts.pop()
    title, tables = None, []
    for loc, p in files:
        t, rows = read_rows(p)
        check_readable(loc, p, rows)
        title = title or t
        tables.append((loc, rows))
    header, rows, counts, dropped = merge_tables(tables)
    for loc, p in parts:
        counts.setdefault(loc, 0)
    total = len(rows)
    if len(files) == 1:
        return files[0][1], counts, total, dropped
    out_path = os.path.splitext(out_path)[0] + ext
    if ext == ".csv":
        write_csv(out_path, header, rows)
    else:
        write_xlsx(out_path, title, header, rows)
    return out_path, counts, total, dropped


def log_merge(name, counts, total, dropped, order=None):
    order = order or list(counts)
    print(f"\n    {name}: rows per location -> "
          + ", ".join(f"{loc}={counts.get(loc, 0)}" for loc in order)
          + f" | total {total}"
          + (f" ({dropped} duplicate row(s) seen in an earlier location dropped)" if dropped else ""))


def check_readable(loc, path, rows):
    """Every Vinculum report has many columns: a file read as ONE column (or none) was
    misread (the 2026-10-21 Order Enquiry regression) - refuse rather than upload it."""
    rows = [_trim(r) for r in rows or []]
    while rows and not rows[0]:
        rows.pop(0)
    if not rows:
        return
    named = [h for h in rows[0] if not _blank(h)]
    if len(named) < 2:
        raise MergeError(f"{loc}: only {len(named)} column(s) read from {os.path.basename(path)} "
                         f"(header {rows[0]!r}) - the file was misread; not merging it.")


def check_columns(path, required):
    """MergeError unless the file's header has every required column (as the
    downstream readers expect them)."""
    _, rows = read_rows(path)
    header = [("" if h is None else str(h).strip()) for h in (rows[0] if rows else [])]
    missing = [c for c in required if c not in header]
    if missing:
        raise MergeError(f"{os.path.basename(path)} lacks column(s) {missing} "
                         f"(has {len(header)} columns) - not uploading it.")


def merge_for_upload(parts, work_dir, drive_name, required_columns=None):
    """Merges the per-location files for one Drive file and logs the per-location
    row counts. Returns (path_to_upload or None, total_rows). required_columns:
    column names the uploaded file must have (else MergeError)."""
    safe = "".join(c if c.isalnum() or c in "._-" else "_" for c in drive_name)
    out = os.path.join(work_dir, f"merged_{safe}")
    path, counts, total, dropped = merge_location_files(parts, out)
    log_merge(drive_name, counts, total, dropped, order=[loc for loc, _ in parts])
    if path and required_columns:
        check_columns(path, required_columns)
    return path, total


def dry_run():
    """VIN_DRY_RUN=true: export + merge as usual, but do NOT upload to Drive."""
    return (os.environ.get("VIN_DRY_RUN") or "").strip().lower() in ("1", "true", "yes", "y")


def keep_dry_run_copy(path, debug_dir, drive_name):
    import shutil
    os.makedirs(debug_dir, exist_ok=True)
    dest = os.path.join(debug_dir, f"DRYRUN_{drive_name}{os.path.splitext(path)[1]}")
    shutil.copyfile(path, dest)
    print(f"    [dry run] NOT uploaded to Drive; merged file kept as {dest}")
    return dest


def remove_files(paths):
    for p in {p for p in paths if p}:
        try:
            os.remove(p)
        except OSError:
            pass


def remove_attempt_files(parts, merged, keep_parts=False):
    """End of one export attempt: deletes the merged file and the per-location files.
    keep_parts=True (only after a failed Drive UPLOAD, 2026-10-23): the per-location
    files stay on disk, so the retry reuses them (export_all_locations / export_one_location
    done=...) instead of exporting every location again. After any other failure they
    are deleted as before, so a bad export is never reused."""
    part_paths = [p for _, p in parts if p]
    remove_files(([] if keep_parts else part_paths)
                 + ([merged] if merged and merged not in part_paths else []))
